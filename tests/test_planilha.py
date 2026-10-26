import tempfile
import unittest
from pathlib import Path
import re
import warnings
from zipfile import ZipFile

from openpyxl import Workbook

from fortigate_audit.planilha import PlanilhaWarning, load_targets


class PlanilhaTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.path = Path(self.directory.name) / "inventario.xlsx"
        self.workbook = Workbook()

    def tearDown(self):
        self.workbook.close()
        self.directory.cleanup()

    def save(self):
        self.workbook.save(self.path)

    def test_multi_sheet_unicode_headers_labels_and_duplicates(self):
        first = self.workbook.active
        first.title = "SP"
        first.append(["Inventário das filiais"])
        first.append(["Equipamento", "Endereço IP"])
        first.append(["Loja São Paulo", "203.0.113.5"])
        second = self.workbook.create_sheet("Oculta")
        second.sheet_state = "hidden"
        second.append(["Host_Name", "IP Público"])
        second.append(["fw.example.com", "2001:db8::1"])
        second.append(["Loja São Paulo", "203.0.113.5"])
        self.save()
        targets = load_targets(self.path)
        self.assertEqual(len(targets), 2)
        self.assertEqual((targets[0].host, targets[0].address, targets[0].tls_name), ("Loja São Paulo", "203.0.113.5", "203.0.113.5"))
        self.assertEqual(targets[1].tls_name, "fw.example.com")
        self.assertEqual(targets[0].source, "SP:linha 3")

    def test_hostname_only_and_idna(self):
        sheet = self.workbook.active
        sheet.append(["Hostname"])
        sheet.append(["Caixa.Example.COM."])
        sheet.append(["caixa.exemplo.çom"])
        self.save()
        targets = load_targets(self.path)
        self.assertEqual(targets[0].address, "caixa.example.com")
        self.assertEqual(targets[1].address, "caixa.exemplo.xn--om-3ia")

    def test_inventory_label_is_not_used_as_tls_hostname(self):
        sheet = self.workbook.active
        sheet.append(["Host", "IP"])
        sheet.append(["LOJA001", "192.0.2.1"])
        sheet.append(["fw.example.com", "192.0.2.2"])
        sheet.append(["2001:db8::5", "2001:db8::5"])
        self.save()
        targets = load_targets(self.path)
        self.assertEqual(targets[0].host, "LOJA001")
        self.assertEqual(targets[0].tls_name, "192.0.2.1")
        self.assertEqual(targets[1].tls_name, "fw.example.com")
        self.assertEqual(targets[2].tls_name, "2001:db8::5")

    def test_explicit_column_names_and_selected_sheet(self):
        sheet = self.workbook.active
        sheet.title = "Alvos"
        sheet.append(["Descrição da caixa", "WAN primária"])
        sheet.append(["Filial sul", "192.0.2.8"])
        self.workbook.create_sheet("Outro").append(["Outras informações"])
        self.save()
        targets = load_targets(self.path, sheet="Alvos", host_column="Descrição da caixa", ip_column="WAN primária")
        self.assertEqual(targets[0].address, "192.0.2.8")

    def test_malformed_address_blocks_partial_inventory(self):
        sheet = self.workbook.active
        sheet.append(["Host", "IP"])
        sheet.append(["Bom", "192.0.2.1"])
        sheet.append(["Ruim", "https://example.com:4443/login"])
        self.save()
        with self.assertRaisesRegex(ValueError, "linha 3"):
            load_targets(self.path)

    def test_formula_rejected(self):
        sheet = self.workbook.active
        sheet.append(["Host", "IP"])
        sheet.append(["Filial", '=CONCAT("192.0.2.","1")'])
        self.save()
        with self.assertRaisesRegex(ValueError, "fórmula"):
            load_targets(self.path)

    def test_unrecognized_sheet_is_reported(self):
        sheet = self.workbook.active
        sheet.append(["Host", "IP"])
        sheet.append(["Filial", "192.0.2.1"])
        self.workbook.create_sheet("Notas").append(["Anotações"])
        self.save()
        with warnings.catch_warnings(record=True) as notices:
            warnings.simplefilter("always")
            self.assertEqual(len(load_targets(self.path)), 1)
        self.assertEqual(len(notices), 1)
        self.assertIs(notices[0].category, PlanilhaWarning)
        self.assertIn("Notas", str(notices[0].message))

    def test_ambiguous_headers_fail_until_explicit(self):
        sheet = self.workbook.active
        sheet.append(["Host", "IP Público", "WAN IP"])
        sheet.append(["Filial", "192.0.2.1", "192.0.2.2"])
        self.save()
        with self.assertRaisesRegex(ValueError, "ambíguos"):
            load_targets(self.path)
        targets = load_targets(self.path, ip_column="WAN IP")
        self.assertEqual(targets[0].address, "192.0.2.2")

    def test_header_like_target_does_not_hide_earlier_targets(self):
        sheet = self.workbook.active
        for row in (["Host"], ["fw1.example.com"], ["dns"], ["fw2.example.com"]):
            sheet.append(row)
        self.save()
        with self.assertRaisesRegex(ValueError, "linhas 1, 3.*--coluna-host"):
            load_targets(self.path)
        targets = load_targets(self.path, host_column="Host")
        self.assertEqual(
            [target.address for target in targets],
            ["fw1.example.com", "dns", "fw2.example.com"],
        )

    def test_control_characters_rejected(self):
        sheet = self.workbook.active
        sheet.append(["Host", "IP"])
        sheet.append(["Filial\nEXPOSTO", "192.0.2.1"])
        self.save()
        with self.assertRaisesRegex(ValueError, "controle"):
            load_targets(self.path)

    def test_control_characters_in_sheet_names_rejected(self):
        sheet = self.workbook.active
        sheet.title = "Inventário\nEXPOSTO"
        sheet.append(["Host", "IP"])
        sheet.append(["Filial", "192.0.2.1"])
        self.save()
        with self.assertRaisesRegex(ValueError, "controle"):
            load_targets(self.path)

    def test_empty_inventory_and_numeric_address_rejected(self):
        sheet = self.workbook.active
        sheet.append(["Host", "IP"])
        self.save()
        with self.assertRaisesRegex(ValueError, "nenhum"):
            load_targets(self.path)
        sheet.append(["Filial", 192])
        self.save()
        with self.assertRaisesRegex(ValueError, "texto"):
            load_targets(self.path)

    def test_corrupt_xlsx_and_invalid_worksheet_xml_have_clear_errors(self):
        self.path.write_bytes(b"isto nao e um arquivo XLSX")
        with self.assertRaisesRegex(ValueError, "corrompido.*XLSX"):
            load_targets(self.path)

        sheet = self.workbook.active
        sheet.append(["Host", "IP"])
        sheet.append(["Filial", "192.0.2.1"])
        self.save()
        with ZipFile(self.path) as archive:
            contents = {name: archive.read(name) for name in archive.namelist()}
        contents["xl/worksheets/sheet1.xml"] = b"<worksheet><sheetData>"
        with ZipFile(self.path, "w") as archive:
            for name, content in contents.items():
                archive.writestr(name, content)
        with self.assertRaisesRegex(ValueError, "corrompido.*XLSX"):
            load_targets(self.path)

    def test_understated_dimensions_do_not_drop_rows_or_columns(self):
        sheet = self.workbook.active
        for _ in range(19):
            sheet.append([])
        sheet.append([None, None, "Host", "IP"])
        sheet.append([None, None, "Filial A", "192.0.2.1"])
        sheet.append([None, None, "Filial B", "192.0.2.2"])
        self.save()
        with ZipFile(self.path) as archive:
            contents = {name: archive.read(name) for name in archive.namelist()}
        contents["xl/worksheets/sheet1.xml"] = re.sub(
            rb'<dimension ref="[^"]+"',
            b'<dimension ref="A1:A2"',
            contents["xl/worksheets/sheet1.xml"],
        )
        with ZipFile(self.path, "w") as archive:
            for name, content in contents.items():
                archive.writestr(name, content)
        targets = load_targets(self.path)
        self.assertEqual([target.address for target in targets], ["192.0.2.1", "192.0.2.2"])
        self.assertEqual(targets[0].source, "Sheet:linha 21")


if __name__ == "__main__":
    unittest.main()
