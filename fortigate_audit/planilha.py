"""Leitura de inventários XLSX sem executar fórmulas nem omitir erros de linha."""

from contextlib import contextmanager
import ipaddress
from pathlib import Path
import re
import unicodedata
import warnings
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

try:
    from lxml.etree import XMLSyntaxError
except ImportError:
    XML_ERRORS = (ParseError,)
else:
    XML_ERRORS = (ParseError, XMLSyntaxError)

from .models import Target


class PlanilhaWarning(UserWarning):
    """Uma aba não pôde ser reconhecida e não entrou na verificação."""


HOST_HEADERS = {
    "host", "hostname", "nomedohost", "equipamento", "nome", "firewall",
    "fortigate", "caixa", "unidade", "site", "cliente",
}
ADDRESS_HEADERS = {
    "ip", "ippublico", "enderecopublico", "enderecoip", "ipwan", "wanip",
    "ipexterno", "ipv4", "ipv6", "ipaddress", "address", "dns", "fqdn",
}
INVALID_XLSX_ERRORS = (BadZipFile, InvalidFileException) + XML_ERRORS
INVALID_XLSX_MESSAGE = "arquivo corrompido ou não é um XLSX válido; abra e salve novamente como .xlsx"


@contextmanager
def _readable_workbook(path: Path):
    # Mantém o arquivo aberto durante a leitura lazy, fechando-o também quando
    # openpyxl falha antes de devolver um workbook.
    with path.open("rb") as source:
        try:
            workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        except INVALID_XLSX_ERRORS as exc:
            raise ValueError(INVALID_XLSX_MESSAGE) from exc
        except KeyError as exc:
            # ZIP sem os componentes obrigatórios de um XLSX.
            raise ValueError(INVALID_XLSX_MESSAGE) from exc
        try:
            yield workbook
        except INVALID_XLSX_ERRORS as exc:
            # read_only lê o XML das linhas sob demanda.
            raise ValueError(INVALID_XLSX_MESSAGE) from exc
        finally:
            workbook.close()


def _header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(c.lower() for c in text if c.isalnum() and not unicodedata.combining(c))


def _text(value: object, field: str) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ValueError(f"{field} deve conter texto, não um valor numérico ou uma data")
    if any(unicodedata.category(c).startswith("C") for c in value):
        raise ValueError(f"{field} contém caracteres de controle")
    return value.strip()


def _address(value: str) -> str:
    """Aceita somente IP literal ou nome DNS, sem porta, URL ou caminho."""
    if not value or len(value) > 253 or any(c.isspace() for c in value):
        raise ValueError("informe um IP ou nome DNS válido, sem espaços")
    if any(c in value for c in "/\\@?#%"):
        raise ValueError("informe somente o IP ou DNS, sem URL, caminho ou escopo")
    try:
        return str(ipaddress.ip_address(value))
    except ValueError:
        pass
    if ":" in value or re.fullmatch(r"[0-9.]+", value):
        raise ValueError("IP inválido ou endereço acompanhado de porta")
    try:
        dns = value.rstrip(".").encode("idna").decode("ascii").lower()
    except UnicodeError as exc:
        raise ValueError("nome DNS inválido") from exc
    if len(dns) > 253 or any(
        not re.fullmatch(r"[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?", label)
        for label in dns.split(".")
    ):
        raise ValueError("nome DNS inválido")
    return dns


def _choose_header(rows: list[tuple], host_column: str | None, ip_column: str | None):
    host_names = {_header(host_column)} if host_column else HOST_HEADERS
    ip_names = {_header(ip_column)} if ip_column else ADDRESS_HEADERS
    candidates = []
    for row_number, row in enumerate(rows, 1):
        hosts = [i for i, value in enumerate(row) if _header(value) in host_names]
        addresses = [i for i, value in enumerate(row) if _header(value) in ip_names]
        if host_column and not hosts or ip_column and not addresses:
            continue
        if hosts or addresses:
            candidates.append((row_number, hosts, addresses, row))
    if not candidates:
        return None
    if len(candidates) > 1:
        row_numbers = ", ".join(str(candidate[0]) for candidate in candidates)
        raise ValueError(
            f"cabeçalho ambíguo: as linhas {row_numbers} parecem cabeçalhos. "
            "Defina os nomes com --coluna-host e/ou --coluna-ip; "
            "se necessário, ajuste os nomes na planilha para evitar ambiguidade"
        )
    row_number, hosts, addresses, row = candidates[0]
    if len(hosts) > 1 or len(addresses) > 1:
        fields = ", ".join(repr(row[i]) for i in sorted(set(hosts + addresses)))
        raise ValueError(
            f"cabeçalhos ambíguos na linha {row_number}: {fields}. "
            "Escolha os nomes com --coluna-host e --coluna-ip"
        )
    return row_number, hosts[0] if hosts else None, addresses[0] if addresses else None


def load_targets(
    path: Path,
    sheet: str | None = None,
    host_column: str | None = None,
    ip_column: str | None = None,
) -> list[Target]:
    """Lê todas as abas, inclusive ocultas, e rejeita linhas de alvo malformadas.

    Cabeçalhos são procurados nas primeiras 20 linhas de cada aba. Abas vazias
    são ignoradas; abas não reconhecidas geram PlanilhaWarning. Fórmulas em
    colunas de alvo não são aceitas, mesmo quando possuem valor em cache.
    Um host com apenas um rótulo (ex.: LOJA001) é tratado como identificação
    do inventário quando há IP/DNS separado: o endereço será usado no TLS.
    Um FQDN ou IP no campo host pode ser usado para validar o certificado.
    """
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("a planilha deve estar no formato .xlsx")
    if host_column is not None and not _header(host_column):
        raise ValueError("--coluna-host não pode ser vazio")
    if ip_column is not None and not _header(ip_column):
        raise ValueError("--coluna-ip não pode ser vazio")
    targets: list[Target] = []
    seen: set[tuple[str, str, str]] = set()
    errors = []
    with _readable_workbook(path) as workbook:
        if sheet is not None and sheet not in workbook.sheetnames:
            names = ", ".join(repr(name) for name in workbook.sheetnames)
            raise ValueError(f"aba {sheet!r} não encontrada; disponíveis: {names}")
        sheets = [workbook[sheet]] if sheet is not None else workbook.worksheets
        for worksheet in sheets:
            # Alguns exportadores gravam dimensões menores que os dados reais.
            # Em read_only isso truncaria linhas/colunas sem qualquer aviso.
            worksheet.reset_dimensions()
            try:
                _text(worksheet.title, "nome da aba")
            except ValueError as exc:
                errors.append(f"Aba {worksheet.title!r}: {exc}")
                continue
            preview = list(worksheet.iter_rows(min_row=1, max_row=20, values_only=True))
            try:
                header = _choose_header(preview, host_column, ip_column)
            except ValueError as exc:
                errors.append(f"Aba {worksheet.title!r}: {exc}")
                continue
            if header is None:
                nonempty = any(
                    any(value is not None and str(value).strip() for value in row)
                    for row in worksheet.iter_rows(values_only=True)
                )
                if nonempty:
                    message = (
                        f"Aba {worksheet.title!r} ignorada: cabeçalhos de host/IP não reconhecidos "
                        "nas primeiras 20 linhas. Use --aba, --coluna-host e --coluna-ip"
                    )
                    if sheet is not None:
                        errors.append(message)
                    else:
                        warnings.warn(message, PlanilhaWarning, stacklevel=2)
                continue
            header_row, host_index, ip_index = header
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=header_row + 1), header_row + 1
            ):
                if all(cell.value is None or cell.value == "" for cell in row):
                    continue
                source = f"{worksheet.title}:linha {row_number}"
                try:
                    selected = [i for i in (host_index, ip_index) if i is not None]
                    if any(i < len(row) and row[i].data_type == "f" for i in selected):
                        raise ValueError("fórmula em coluna de host/IP; substitua por um valor de texto")
                    host = _text(row[host_index].value if host_index is not None and host_index < len(row) else None, "host")
                    address_text = _text(row[ip_index].value if ip_index is not None and ip_index < len(row) else None, "IP/DNS")
                    if not host and not address_text:
                        raise ValueError("linha preenchida sem host ou IP/DNS")
                    address = _address(address_text or host)
                    host = host or address
                    try:
                        tls_name = _address(host)
                        if address_text and "." not in tls_name:
                            try:
                                ipaddress.ip_address(tls_name)
                            except ValueError:
                                tls_name = address
                    except ValueError:
                        tls_name = address
                    key = (host.casefold(), address, tls_name)
                    if key not in seen:
                        targets.append(Target(host, address, tls_name, source))
                        seen.add(key)
                except ValueError as exc:
                    errors.append(f"{source}: {exc}")
        if errors:
            detail = "\n".join(errors[:10])
            if len(errors) > 10:
                detail += f"\n... e mais {len(errors) - 10} erro(s)"
            raise ValueError("Corrija a planilha antes de verificar:\n" + detail)
        if not targets:
            raise ValueError("nenhum host/IP encontrado; confira os cabeçalhos e a aba selecionada")
        return targets
